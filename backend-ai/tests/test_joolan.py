import io
import json
import re
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from types import ModuleType
from openpyxl import Workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import joolan
from joolan_client import Client, ExportError


def row(**changes):
    data=dict(Order_Number='001',Customer='0007',Order_Date='2026-01-02',Product='000123',Ordered_Quantity=5.5,
              Delivered_Quantity=2.25,Store='A',Colour='01',Size='08',Document_Type='ORDER',Link_Status='Linked item')
    data.update(changes);return data


def workbook(rows,partial=False):
    wb=Workbook();wb.active.title='negoce_entetes';wb.active.append(['Entete'])
    ws=wb.create_sheet('Trade documents');fields=joolan.FIELDS+joolan.EXTRA;ws.append(fields)
    for data in rows:ws.append([data.get(c) for c in fields])
    if partial:
        details=wb.create_sheet('Export details');details.append(['negoce_lignes',None,None,None,None,None,'INCOMPLETE'])
    out=io.BytesIO();wb.save(out);return out.getvalue()


class SnapshotTests(unittest.TestCase):
    def setUp(self):
        self.conn=sqlite3.connect(':memory:')
        joolan.ensure_joolan_schema(self.conn)
        self.conn.execute("INSERT INTO joolan_state VALUES('customer_order_type','ORDER')")
        self.conn.commit()
    def tearDown(self):self.conn.close()
    def test_selects_trade_sheet_not_first_raw_sheet(self):
        result=joolan.replace_snapshot(self.conn,joolan.read_trade_workbook(workbook([row()])),'workbook')
        self.assertEqual(result['usableRows'],1)
        saved=joolan.snapshot_rows(self.conn)[0]
        self.assertEqual(saved['Product'],'000123')
        self.assertEqual(saved['Ordered_Quantity'],5.5)
    def test_reimport_is_idempotent_and_changed_snapshot_replaces(self):
        joolan.replace_snapshot(self.conn,[row(),row(Order_Number='002')],'file')
        again=joolan.replace_snapshot(self.conn,[row(),row(Order_Number='002')],'file')
        self.assertEqual(again['rowsSkipped'],2)
        self.assertEqual(again['totalRows'],2)
        changed=joolan.replace_snapshot(self.conn,[row(Ordered_Quantity=8)],'file')
        self.assertEqual(changed['rowsUpdated'],1);self.assertEqual(changed['rowsRemoved'],1)
        self.assertEqual(joolan.pipeline_dataframe(self.conn)['CDE_qte'].sum(),8)
    def test_identical_item_lines_are_not_deduped(self):
        result=joolan.replace_snapshot(self.conn,[row(),row()],'file')
        self.assertEqual(result['usableRows'],2)
        self.assertEqual(joolan.pipeline_dataframe(self.conn)['CDE_qte'].sum(),11)
    def test_invalid_rows_remain_for_review_but_are_not_forecast(self):
        result=joolan.replace_snapshot(self.conn,[row(),row(Order_Number='002',Order_Date=None),row(Order_Number='003',Link_Status='No unique compatible header'),row(Order_Number='004',Delivered_Quantity=10)],'file')
        self.assertEqual(result['quarantinedRows'],3)
        self.assertEqual(len(joolan.snapshot_rows(self.conn)),4)
        self.assertEqual(len(joolan.pipeline_dataframe(self.conn)),1)
    def test_incomplete_or_all_invalid_snapshot_preserves_previous(self):
        joolan.replace_snapshot(self.conn,[row()],'file')
        with self.assertRaisesRegex(ValueError,'incomplete'):
            joolan.replace_snapshot(self.conn,joolan.read_trade_workbook(workbook([row()],True)),'file')
        with self.assertRaisesRegex(ValueError,'previous'):
            joolan.replace_snapshot(self.conn,[row(Customer=None)],'file')
        self.assertEqual(joolan.snapshot_summary(self.conn)['totalRows'],1)
    def test_incomplete_export_reports_missing_source_counts(self):
        wb = Workbook()
        ws = wb.active; ws.title = 'Export details'
        ws.append(['Dataset', 'Exported rows', 'Expected rows', 'Status'])
        ws.append(['negoce_entetes', 244924, 244924, 'COMPLETE'])
        ws.append(['negoce_lignes', 566000, 2741625, 'INCOMPLETE'])
        out = io.BytesIO(); wb.save(out)
        with self.assertRaisesRegex(ValueError, r'negoce_lignes \(566,000 of 2,741,625 rows\)'):
            list(joolan.read_trade_workbook(out.getvalue()))
    def test_pipeline_source_is_explicit_and_preserves_ids_units(self):
        from m1_data_preparation import load_data,parse_and_prepare,aggregate_monthly_sku_color_size
        joolan.replace_snapshot(self.conn,[row(),row(Order_Number='002')],'file')
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'snapshot.csv'
            self.assertIsNone(joolan.selected_pipeline_path(self.conn,path))
            joolan.select_forecast_source(self.conn,'joolan')
            joolan.selected_pipeline_path(self.conn,path)
            prepared=parse_and_prepare(load_data(str(path)))
            self.assertEqual(prepared['Client'].iloc[0],'0007')
            self.assertEqual(prepared['Reference'].iloc[0],'000123')
            self.assertEqual(prepared['CDE_qte'].sum(),11)
            self.assertEqual(prepared['true_demand'].sum(),11)
            monthly=aggregate_monthly_sku_color_size(prepared)
            self.assertEqual(monthly['true_demand'].sum(),11)
            self.assertEqual(monthly['CDE_qte'].sum(),11)
    def test_selection_survives_pipeline_generated_table_rebuild(self):
        import database
        database._create_persistent_tables(self.conn);database.ensure_pipeline_schema(self.conn)
        joolan.replace_snapshot(self.conn,[row()],'file');joolan.select_forecast_source(self.conn,'joolan')
        database._create_schema(self.conn)
        self.assertEqual(joolan.snapshot_summary(self.conn)['forecastSource'],'joolan')
        self.assertEqual(joolan.snapshot_summary(self.conn)['usableRows'],1)


class ApiTests(unittest.TestCase):
    def test_file_api_preview_pagination_and_source_selection(self):
        import api_server as api
        from fastapi.testclient import TestClient
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'api.db'
            def connection():return sqlite3.connect(path)
            with patch.object(api,'get_connection',connection),TestClient(api.app) as client:
                # Legacy importer regression fixture; fresh production rejects uploads.
                conn=connection();conn.execute("DELETE FROM joolan_state WHERE key='source_policy'");conn.commit();conn.close()
                result=client.post('/api/joolan/file',files={'file':('Joolan Trade Documents.xlsx',workbook([row(),row(Order_Number='002')]),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
                self.assertEqual(result.status_code,200,result.text)
                self.assertEqual(result.json()['snapshot']['usableRows'],2)
                self.assertEqual(client.get('/api/joolan/orders?offset=1&limit=1').json()['rows'][0]['Order_Number'],'002')
                self.assertEqual(client.post('/api/joolan/order-type',json={'documentType':'ORDER'}).status_code,200)
                self.assertEqual(client.post('/api/joolan/forecast-source',json={'source':'joolan'}).json()['forecastSource'],'joolan')
                failed=client.post('/api/joolan/file',files={'file':('bad.xlsx',b'not a workbook')})
                self.assertEqual(failed.status_code,400)
                self.assertEqual(client.get('/api/joolan/summary').json()['totalRows'],2)
                generic=client.post('/api/ingest/file',files={'file':('trade.xlsx',workbook([row()]),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
                self.assertEqual(generic.status_code,200,generic.text)
                self.assertEqual(generic.json()['source'],'joolan')
    def test_forecast_job_receives_only_selected_joolan_snapshot(self):
        import api_server as api
        import database
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'forecast.db'
            def connection():return sqlite3.connect(path)
            conn=connection();database._create_persistent_tables(conn)
            joolan.replace_snapshot(conn,[row(),row(Order_Number='002')],'file')
            joolan.confirm_order_type(conn,'ORDER')
            joolan.select_forecast_source(conn,'joolan')
            api.insert_job_running(conn,'forecast1','forecast');conn.commit();conn.close()
            fake=ModuleType('main');fake.run_phase1=Mock()
            with patch.object(api,'get_connection',connection),patch.object(api,'BASE_DIR',Path(d)),patch.dict(sys.modules,{'main':fake}):
                api._run_forecast_pipeline_job('forecast1')
                args=fake.run_phase1.call_args.kwargs
                self.assertIn('.joolan_pipeline.csv',args['data_path'])
                import pandas as pd
                data=pd.read_csv(args['data_path'],dtype=str)
                self.assertEqual(len(data),2)
                self.assertTrue(data['Data_Source'].eq('joolan').all())
                self.assertEqual(api.get_ingest_status('forecast1')['status'],'complete')
    def test_api_sync_uses_both_tables_and_line_quantities(self):
        sources={'negoce_entetes':[{'Entete':1,'Client':'007','Date':'2026-01-02','Magasin':'A','Nature':'ORDER'}],
                 'negoce_lignes':[{'Entete':1,'Client':'007','Magasin':'A','Produit':'00123','Quantite':5.5,'Quantite_Livree':1.5},
                                  {'Entete':1,'Client':'007','Magasin':'A','Produit':'00234','Quantite':2,'Quantite_Livree':0}]}
        calls=[]
        class Fake(Client):
            def query(self,sql):
                calls.append(sql)
                if 'information_schema' in sql:return []
                table=next(t for t in sources if '`'+t+'`' in sql)
                if 'COUNT(*)' in sql:return [{'export_count':len(sources[table])}]
                limit=re.search(r'LIMIT (\d+)',sql);offset=re.search(r'OFFSET (\d+)',sql)
                start=int(offset[1]) if offset else 0
                return sources[table][start:start+min(int(limit[1]),1)]
        with patch.object(joolan,'Client',Fake):
            rows=list(joolan.fetch_trade_rows('not-real','legend73','ORDER',1000))
        self.assertEqual(len(rows),2);self.assertEqual(sum(r['Ordered_Quantity'] for r in rows),7.5)
        self.assertTrue(any('OFFSET 1' in q for q in calls))
    def test_failed_sync_preserves_snapshot_and_hides_key(self):
        import api_server as api
        import database
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'sync.db'
            def connection():return sqlite3.connect(path)
            conn=connection();database._create_persistent_tables(conn)
            joolan.replace_snapshot(conn,[row()],'file')
            api.insert_job_running(conn,'job1','joolan-api');conn.commit();conn.close()
            with patch.object(api,'get_connection',connection),patch.object(api,'fetch_trade_rows',side_effect=ValueError('bad key SECRET')):
                api._run_joolan_sync('job1','SECRET','legend73','',1000)
                status=api.get_ingest_status('job1')
                self.assertEqual(status['status'],'error');self.assertNotIn('SECRET',status['errorMessage'])
                self.assertEqual(api.get_joolan_summary()['totalRows'],1)

if __name__=='__main__':unittest.main()
